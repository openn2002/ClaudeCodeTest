"""
Digital Wellness — Enterprise Proposal Deck Builder
====================================================
Generates a branded PowerPoint presentation using the Digital Wellness /
CSIRO Total Wellbeing Diet visual identity.

Usage:
    python build_deck.py

Output:
    reports/enterprise-proposal-FY26.pptx
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import datetime

# ── Brand colour palette ────────────────────────────────────────────────────
TEAL        = RGBColor(0x00, 0xC4, 0xCC)   # DW / CSIRO TWD primary teal
TEAL_DARK   = RGBColor(0x00, 0x9A, 0xA0)   # darker teal for contrast
NAVY        = RGBColor(0x1A, 0x2B, 0x47)   # dark navy (footer, dark slides)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
DARK_TEXT   = RGBColor(0x1F, 0x29, 0x37)   # near-black headings
BODY_TEXT   = RGBColor(0x44, 0x44, 0x44)   # dark grey body
MID_GREY    = RGBColor(0x88, 0x88, 0x88)   # secondary / captions
LIGHT_BG    = RGBColor(0xF0, 0xF7, 0xF8)   # card / table row bg
RULE_GREY   = RGBColor(0xCC, 0xD6, 0xDC)   # horizontal rules

SLIDE_W = Inches(13.33)
SLIDE_H = Inches(7.5)
FOOTER_H = 0.42
FOOTER_Y = 7.5 - FOOTER_H

prs = Presentation()
prs.slide_width  = SLIDE_W
prs.slide_height = SLIDE_H

BLANK = prs.slide_layouts[6]  # completely blank


# ── Low-level helpers ────────────────────────────────────────────────────────

def rect(slide, l, t, w, h, fill=None, line=False):
    """Add a rectangle. l/t/w/h in inches."""
    shp = slide.shapes.add_shape(
        1, Inches(l), Inches(t), Inches(w), Inches(h))
    shp.line.fill.background() if not line else None
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    else:
        shp.fill.background()
    return shp


def txt(slide, text, l, t, w, h,
        size=14, bold=False, italic=False,
        color=DARK_TEXT, align=PP_ALIGN.LEFT, wrap=True):
    """Add a text box. l/t/w/h in inches."""
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf  = box.text_frame
    tf.word_wrap = wrap
    p   = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size    = Pt(size)
    run.font.bold    = bold
    run.font.italic  = italic
    run.font.color.rgb = color
    return box


def dot_pattern(slide, l, t, cols=6, rows=5, spacing=0.18, r=0.05,
                color=TEAL, alpha=60):
    """Draw a grid of small circles as a decorative dot pattern."""
    for row in range(rows):
        for col in range(cols):
            x = l + col * spacing
            y = t + row * spacing
            shp = slide.shapes.add_shape(
                9,  # oval
                Inches(x), Inches(y), Inches(r * 2), Inches(r * 2))
            shp.fill.solid()
            shp.fill.fore_color.rgb = color
            shp.line.fill.background()


# ── Brand building blocks ────────────────────────────────────────────────────

def footer(slide, page_num, url="www.digitalwellness.com"):
    """Standard dark-navy footer bar with DW logo text, URL, page number."""
    rect(slide, 0, FOOTER_Y, 13.33, FOOTER_H, fill=NAVY)
    # DW wordmark (text approximation — replace with image when logo asset available)
    txt(slide, "Digital Wellness\u00ae",
        0.25, FOOTER_Y + 0.05, 2.2, 0.32,
        size=9, bold=True, color=WHITE)
    txt(slide, f"|   {url}",
        2.5, FOOTER_Y + 0.06, 5.0, 0.30,
        size=8, color=MID_GREY)
    txt(slide, f"PAGE {page_num}",
        11.8, FOOTER_Y + 0.06, 1.3, 0.30,
        size=8, color=MID_GREY, align=PP_ALIGN.RIGHT)


def white_bg(slide):
    rect(slide, 0, 0, 13.33, 7.5, fill=WHITE)


def teal_bg(slide):
    rect(slide, 0, 0, 13.33, 7.5, fill=TEAL)


def navy_bg(slide):
    rect(slide, 0, 0, 13.33, 7.5, fill=NAVY)


def photo_placeholder(slide, l, t, w, h, label="[ Photo ]"):
    """Light teal placeholder for image areas."""
    rect(slide, l, t, w, h, fill=LIGHT_BG)
    txt(slide, label, l + w / 2 - 0.6, t + h / 2 - 0.2, 1.2, 0.4,
        size=9, color=MID_GREY, align=PP_ALIGN.CENTER)


def toc_row(slide, num, title, y):
    """One row of a Table of Contents."""
    txt(slide, num, 0.55, y, 0.55, 0.38,
        size=16, bold=True, color=TEAL)
    txt(slide, title, 1.15, y + 0.04, 6.5, 0.35,
        size=13, color=DARK_TEXT)
    rect(slide, 0.55, y + 0.42, 7.1, 0.02, fill=RULE_GREY)


def step_card(slide, l, t, w, num_label, title, body):
    """One card in a 5-step process row."""
    rect(slide, l, t, w, 2.2, fill=LIGHT_BG)
    rect(slide, l, t, w, 0.38, fill=TEAL)
    txt(slide, num_label, l + 0.12, t + 0.06, w - 0.24, 0.28,
        size=10, bold=True, color=WHITE)
    txt(slide, title, l + 0.12, t + 0.46, w - 0.24, 0.36,
        size=11, bold=True, color=DARK_TEXT)
    txt(slide, body, l + 0.12, t + 0.86, w - 0.24, 1.25,
        size=9, color=BODY_TEXT)


def pillar_card(slide, l, t, w, h, icon_char, title, body):
    """One card in a 2×4 pillars grid."""
    rect(slide, l, t, w, h, fill=LIGHT_BG)
    rect(slide, l, t, 0.08, h, fill=TEAL)
    txt(slide, icon_char, l + 0.18, t + 0.10, 0.5, 0.4,
        size=16, color=TEAL)
    txt(slide, title, l + 0.18, t + 0.52, w - 0.3, 0.35,
        size=10, bold=True, color=DARK_TEXT)
    txt(slide, body, l + 0.18, t + 0.90, w - 0.3, h - 1.05,
        size=9, color=BODY_TEXT)


def stat_block(slide, l, t, value, label, value_color=TEAL):
    """A single stat (large number + label)."""
    txt(slide, value, l, t, 2.8, 0.65,
        size=32, bold=True, color=value_color, align=PP_ALIGN.CENTER)
    txt(slide, label, l, t + 0.65, 2.8, 0.38,
        size=10, color=WHITE if value_color == WHITE else BODY_TEXT,
        align=PP_ALIGN.CENTER)


def section_divider_white(slide, page, logo_line, title, subtitle):
    """
    White-background section divider.
    Left: logo line + large bold title + teal subtitle.
    Right: dark navy decorative arch + photo placeholder.
    """
    white_bg(slide)

    # right-side dark navy decorative half-circle
    rect(slide, 7.8, -1.5, 7.0, 10.0, fill=NAVY)
    shp = slide.shapes.add_shape(
        9, Inches(6.8), Inches(-0.5), Inches(7.5), Inches(9.0))
    shp.fill.solid()
    shp.fill.fore_color.rgb = NAVY
    shp.line.fill.background()
    photo_placeholder(slide, 8.4, 1.2, 4.2, 4.8, "[ Image ]")

    # decorative dots top-right (white dots on navy)
    dot_pattern(slide, 12.2, 0.2, cols=4, rows=4,
                spacing=0.22, r=0.04, color=WHITE)

    txt(slide, logo_line,
        0.6, 1.5, 6.0, 0.45,
        size=11, bold=True, color=TEAL)
    txt(slide, title,
        0.6, 2.1, 6.5, 1.8,
        size=36, bold=True, color=DARK_TEXT)
    txt(slide, subtitle,
        0.6, 4.05, 6.5, 0.55,
        size=15, color=TEAL, italic=False)

    footer(slide, page)


def section_divider_teal(slide, page, logo_line, title, subtitle):
    """
    Teal-background section divider (used for CSIRO TWD programme slides).
    """
    teal_bg(slide)
    dot_pattern(slide, 0.3, 0.3, cols=5, rows=4, spacing=0.22,
                r=0.04, color=WHITE)
    dot_pattern(slide, 10.8, 5.2, cols=5, rows=4, spacing=0.22,
                r=0.04, color=WHITE)

    txt(slide, logo_line,
        0.7, 1.4, 7.0, 0.5,
        size=13, bold=True, color=WHITE)
    txt(slide, title,
        0.7, 2.1, 8.0, 1.8,
        size=40, bold=True, color=WHITE)
    txt(slide, subtitle,
        0.7, 4.05, 8.0, 0.5,
        size=15, color=WHITE)

    photo_placeholder(slide, 8.8, 1.0, 4.0, 5.0, "[ Image ]")
    footer(slide, page)


# ════════════════════════════════════════════════════════════════════════════
# SLIDES
# ════════════════════════════════════════════════════════════════════════════

# ── SLIDE 1 — COVER ─────────────────────────────────────────────────────────
slide = prs.slides.add_slide(BLANK)
white_bg(slide)

# teal accent strip left edge
rect(slide, 0, 0, 0.18, 7.5, fill=TEAL)

# decorative dots (bottom-left, top-right)
dot_pattern(slide, 0.35, 5.4, cols=6, rows=4, spacing=0.20, r=0.04, color=TEAL)
dot_pattern(slide, 10.8, 0.25, cols=5, rows=3, spacing=0.22, r=0.04, color=TEAL)

# logo area: DW + CSIRO TWD
txt(slide, "Digital Wellness\u00ae",
    0.55, 0.45, 3.5, 0.5,
    size=13, bold=True, color=NAVY)
txt(slide, "CSIRO  TOTAL WELLBEING DIET",
    4.2, 0.45, 4.5, 0.5,
    size=10, bold=True, color=TEAL)

# title block
txt(slide, "Enterprise Proposal",
    0.55, 1.55, 7.5, 1.0,
    size=42, bold=True, color=DARK_TEXT)
txt(slide, "Partnering to transform weight management\nand deliver a healthier tomorrow for Australia.",
    0.55, 2.75, 7.2, 1.1,
    size=16, color=BODY_TEXT)

rect(slide, 0.55, 4.05, 5.5, 0.04, fill=TEAL)

txt(slide, f"Prepared for:  [Partner Name]",
    0.55, 4.2, 5.5, 0.4,
    size=12, color=BODY_TEXT)
txt(slide, f"Date:  {datetime.date.today().strftime('%B %Y')}",
    0.55, 4.65, 5.5, 0.4,
    size=12, color=MID_GREY)

# right-side photo
photo_placeholder(slide, 8.2, 0.5, 4.8, 6.3, "[ Hero image ]")

footer(slide, 1)


# ── SLIDE 2 — TABLE OF CONTENTS ─────────────────────────────────────────────
slide = prs.slides.add_slide(BLANK)
white_bg(slide)
dot_pattern(slide, 11.8, 0.3, cols=4, rows=5, spacing=0.20, r=0.04, color=TEAL)

txt(slide, "Table of Contents",
    0.55, 0.45, 7.0, 0.65,
    size=26, bold=True, color=DARK_TEXT)

sections = [
    ("01", "About Us"),
    ("02", "The CSIRO Total Wellbeing Diet"),
    ("03", "How We Work With Partners"),
    ("04", "More Information"),
]
for i, (num, title) in enumerate(sections):
    toc_row(slide, num, title, 1.35 + i * 0.9)

photo_placeholder(slide, 8.0, 0.9, 4.8, 5.8, "[ Photo ]")
footer(slide, 2)


# ── SLIDE 3 — SECTION DIVIDER: ABOUT US ────────────────────────────────────
slide = prs.slides.add_slide(BLANK)
section_divider_white(
    slide, 3,
    "Digital Wellness\u00ae",
    "About Us",
    "Making better health possible — for everyone, for life.")


# ── SLIDE 4 — WE KNOW NUTRITION ─────────────────────────────────────────────
slide = prs.slides.add_slide(BLANK)
white_bg(slide)

txt(slide, "We know nutrition, and we know health.",
    0.55, 0.4, 9.0, 0.65,
    size=24, bold=True, color=DARK_TEXT)
txt(slide, "It's in our DNA.",
    0.55, 1.05, 9.0, 0.45,
    size=16, color=TEAL, bold=False)

body = (
    "Digital Wellness was founded in 2006 by Scott Penn, building on three generations of "
    "innovation in weight management and preventive health. From pioneering Australia's weight "
    "loss market in the 1980s to transforming Weight Watchers AU/NZ into the most successful "
    "global unit in the 1990s, the Penn family has a long-standing legacy of health leadership.\n\n"
    "Today, Digital Wellness powers two of the world's most credible, science-backed digital "
    "health programs — the CSIRO Total Wellbeing Diet in Australia and the Mayo Clinic Diet "
    "in the United States — and has positively impacted more than 3 million lives."
)
txt(slide, body,
    0.55, 1.65, 7.8, 2.8,
    size=11, color=BODY_TEXT)

# three stat callouts
stats = [
    ("3M+", "Lives positively\nimpacted"),
    ("15+", "Years of digital\nhealth innovation"),
    ("1.2M", "Australians\nimpacted"),
]
for i, (val, lbl) in enumerate(stats):
    lx = 0.55 + i * 2.65
    rect(slide, lx, 4.75, 2.4, 1.65, fill=LIGHT_BG)
    txt(slide, val, lx + 0.1, 4.88, 2.2, 0.65,
        size=30, bold=True, color=TEAL, align=PP_ALIGN.CENTER)
    txt(slide, lbl, lx + 0.1, 5.52, 2.2, 0.72,
        size=10, color=BODY_TEXT, align=PP_ALIGN.CENTER)

photo_placeholder(slide, 9.0, 1.0, 3.9, 5.4, "[ Founder / team photo ]")
footer(slide, 4)


# ── SLIDE 5 — WHAT SETS OUR SOLUTIONS APART ─────────────────────────────────
slide = prs.slides.add_slide(BLANK)
white_bg(slide)

txt(slide, "What sets our solutions apart",
    0.55, 0.38, 12.0, 0.6,
    size=24, bold=True, color=DARK_TEXT)
txt(slide, "Uniquely positioned to combine science, behaviour change and real-world data.",
    0.55, 0.98, 12.0, 0.38,
    size=12, color=TEAL)

pillars = [
    ("◎", "Continuous Evaluation\nfrom the CSIRO",
     "Programs regularly reviewed and updated through our partnership with Australia's national science agency."),
    ("◈", "100% Digitally Delivered\nwith a Human Touch",
     "Accessible, affordable health solutions for a range of health concerns, allowing scalability and integration into modern life."),
    ("◉", "Access to Digital Tools\n& Allied Health Support",
     "Range of member-informed, scientifically developed digital tools that ensure accountability, engagement and individualised support."),
    ("⬡", "Behaviour Change\nModelled Program",
     "Focusing on educational tools, enhanced decision-making and mindfulness strategies to improve emotional wellbeing and self-control."),
    ("◍", "Nutritional-Based\nPhilosophy & Approach",
     "High protein, low GI philosophy developed and researched by the CSIRO — proven to increase long-term weight loss success."),
    ("◎", "Peer Community Support\n& Group Coaching",
     "Access to one of Australia's largest online peer-support communities, along with complimentary Group Webinar Coaching Sessions."),
    ("⟳", "Connected Care Model",
     "A digital health program focusing on patient-centric care and bridging the gap between dietitians, RNs, patients and health funds."),
    ("★", "Outcomes Beyond\nClinical Results",
     "Members report better energy, mental clarity, confidence, boosted motivation, better sleep and reduced stress."),
]

card_w, card_h = 3.1, 1.72
for i, (icon, title, desc) in enumerate(pillars):
    col = i % 4
    row = i // 4
    lx = 0.45 + col * 3.2
    ty = 1.55 + row * 1.88
    pillar_card(slide, lx, ty, card_w, card_h, icon, title, desc)

footer(slide, 5)


# ── SLIDE 6 — HOW WE CAN HELP (dark) ────────────────────────────────────────
slide = prs.slides.add_slide(BLANK)
navy_bg(slide)
dot_pattern(slide, 0.3, 0.3, cols=5, rows=4, spacing=0.25,
            r=0.045, color=RGBColor(0x25, 0x3C, 0x5E))
dot_pattern(slide, 10.5, 5.0, cols=5, rows=4, spacing=0.25,
            r=0.045, color=RGBColor(0x25, 0x3C, 0x5E))

txt(slide, "How we can help",
    1.0, 1.0, 11.33, 0.75,
    size=30, bold=False, color=WHITE, align=PP_ALIGN.CENTER)

benefits = [
    ("⬡", "Upstream Risk\nIdentification"),
    ("◎", "Improved Health\nOutcomes"),
    ("$", "Value on Health\nInvestment"),
    ("◉", "Improved Member\nRetention"),
]
for i, (icon, label) in enumerate(benefits):
    lx = 1.5 + i * 2.7
    txt(slide, icon, lx, 2.2, 2.2, 0.8,
        size=36, color=TEAL, align=PP_ALIGN.CENTER)
    txt(slide, label, lx, 3.1, 2.2, 0.9,
        size=13, color=WHITE, align=PP_ALIGN.CENTER)

footer(slide, 6)


# ── SLIDE 7 — SECTION DIVIDER: THE CSIRO TWD ────────────────────────────────
slide = prs.slides.add_slide(BLANK)
section_divider_teal(
    slide, 7,
    "CSIRO  TOTAL WELLBEING DIET",
    "Australia's No.1\nRated Diet",
    "Real members. Real Science. Real Results.")


# ── SLIDE 8 — PROGRAM OVERVIEW & KEY STATS ──────────────────────────────────
slide = prs.slides.add_slide(BLANK)
white_bg(slide)

txt(slide, "At the core of all our solutions",
    0.55, 0.38, 9.0, 0.6,
    size=24, bold=True, color=DARK_TEXT)
txt(slide, "Backed by Australia's national science agency — translating world-leading research into practical, preventative health solutions.",
    0.55, 0.98, 9.5, 0.5,
    size=11, color=TEAL)

# four stat badges
badges = [
    ("5 Star\nRated", "Online program — rated No.1 on Trustpilot and Product Review"),
    (">52%", "Retention rate — industry-leading program completion"),
    ("4.8★", "HCP star rating from members on online support sessions"),
    ("95%", "Program satisfaction rating — overall rating of the online program"),
]
for i, (val, lbl) in enumerate(badges):
    lx = 0.45 + i * 3.22
    rect(slide, lx, 1.65, 3.0, 1.6, fill=LIGHT_BG)
    rect(slide, lx, 1.65, 3.0, 0.06, fill=TEAL)
    txt(slide, val, lx + 0.12, 1.78, 2.76, 0.62,
        size=18, bold=True, color=TEAL, align=PP_ALIGN.CENTER)
    txt(slide, lbl, lx + 0.12, 2.40, 2.76, 0.78,
        size=9, color=BODY_TEXT, align=PP_ALIGN.CENTER)

# what's included
txt(slide, "What's included:",
    0.55, 3.52, 8.0, 0.4,
    size=13, bold=True, color=DARK_TEXT)

inclusions = [
    "Access to the CSIRO Total Wellbeing Diet via 12-Week Plan",
    "My Journey AI Guide — CSIRO-researched, member-informed habit builder",
    "Flexible recipes and menu plans based on the high protein, low GI philosophy",
    "Live and recorded Group Coaching Sessions with Accredited Practising Dietitians",
    "Mental Wellbeing Tools, food & exercise tracking (iOS + Android apps)",
    "Access to Australia's largest online peer-support community (50,000+ members)",
]
for i, item in enumerate(inclusions):
    col = i % 2
    row = i // 2
    lx = 0.55 + col * 6.45
    ty = 3.98 + row * 0.54
    txt(slide, f"\u2713  {item}", lx, ty, 6.1, 0.48,
        size=10, color=BODY_TEXT)

photo_placeholder(slide, 10.2, 1.55, 2.8, 5.1, "[ App / program screenshot ]")
footer(slide, 8)


# ── SLIDE 9 — OUTCOMES DATA ──────────────────────────────────────────────────
slide = prs.slides.add_slide(BLANK)
white_bg(slide)

txt(slide, "Program outcomes",
    0.55, 0.38, 9.0, 0.6,
    size=24, bold=True, color=DARK_TEXT)
txt(slide, "Evidence-based results across weight loss, chronic disease markers and member wellbeing.",
    0.55, 0.98, 9.5, 0.4,
    size=12, color=TEAL)

# TWD outcomes row
txt(slide, "CSIRO Total Wellbeing Diet",
    0.55, 1.58, 6.5, 0.4,
    size=13, bold=True, color=DARK_TEXT)
rect(slide, 0.55, 1.98, 12.3, 0.03, fill=TEAL)

twd_stats = [
    ("7.2%", "Average body weight\nloss after 12 weeks"),
    ("1 in 2", "Completers lose 5%+\n(clinically significant)"),
    ("72", "NPS score"),
    ("95%", "Program satisfaction\nrating"),
    ("50K+", "Active community\nmembers"),
]
for i, (val, lbl) in enumerate(twd_stats):
    lx = 0.45 + i * 2.52
    rect(slide, lx, 2.1, 2.35, 1.55, fill=LIGHT_BG)
    txt(slide, val, lx + 0.1, 2.2, 2.15, 0.62,
        size=24, bold=True, color=TEAL, align=PP_ALIGN.CENTER)
    txt(slide, lbl, lx + 0.1, 2.82, 2.15, 0.75,
        size=9, color=BODY_TEXT, align=PP_ALIGN.CENTER)

# CDMP outcomes row
txt(slide, "Total Wellbeing Lifestyle Plan (CDMP)",
    0.55, 3.9, 8.0, 0.4,
    size=13, bold=True, color=DARK_TEXT)
rect(slide, 0.55, 4.3, 12.3, 0.03, fill=TEAL)

cdmp_stats = [
    ("7.1%", "Body weight\nloss"),
    ("2–3%", "Blood pressure\nimprovement"),
    ("5%", "Cholesterol\nimprovement"),
    ("41%", "Reduction in\npressure on joints"),
    ("13%", "Blood glucose\nimprovement"),
]
for i, (val, lbl) in enumerate(cdmp_stats):
    lx = 0.45 + i * 2.52
    rect(slide, lx, 4.42, 2.35, 1.55, fill=LIGHT_BG)
    txt(slide, val, lx + 0.1, 4.52, 2.15, 0.62,
        size=24, bold=True, color=TEAL, align=PP_ALIGN.CENTER)
    txt(slide, lbl, lx + 0.1, 5.14, 2.15, 0.75,
        size=9, color=BODY_TEXT, align=PP_ALIGN.CENTER)

foot_note = ("Source: Hendrie et al. (2021), JMIR 23(6); CSIRO Study — Effectiveness of an online weight "
             "management program for people with chronic disease delivered through a private health "
             "insurance model.")
txt(slide, foot_note,
    0.55, 6.62, 12.3, 0.4,
    size=7, color=MID_GREY, italic=True)

footer(slide, 9)


# ── SLIDE 10 — SECTION DIVIDER: HOW WE WORK WITH PARTNERS ───────────────────
slide = prs.slides.add_slide(BLANK)
section_divider_white(
    slide, 10,
    "Digital Wellness\u00ae",
    "How We Work\nWith Partners",
    "Your success is our success.")


# ── SLIDE 11 — MAXIMISING ROI ───────────────────────────────────────────────
slide = prs.slides.add_slide(BLANK)
white_bg(slide)

txt(slide, "Maximising the return on investment",
    0.55, 0.38, 9.5, 0.65,
    size=24, bold=True, color=DARK_TEXT)
txt(slide, "Getting the right member on the right program at the right time.",
    0.55, 1.02, 9.5, 0.4,
    size=13, color=TEAL)

bullets = [
    ("Right people:",
     "Identifying and investing in those with poor health habits, existing risks, high claims "
     "costs or whose poor health will cost the most in the future."),
    ("Right program:",
     "Matching content, commitment level and anticipated outcomes. We work with each "
     "customer to develop a tailored triage approach, moving members into the right program "
     "at the right time."),
    ("Right time:",
     "Ability to engage members when they are at the right stage of change and ready to "
     "commit to taking action, with year-round program availability."),
]
for i, (bold_part, rest) in enumerate(bullets):
    ty = 1.65 + i * 1.0
    txt(slide, bold_part, 0.55, ty, 1.55, 0.38,
        size=11, bold=True, color=DARK_TEXT)
    txt(slide, rest, 2.15, ty, 5.8, 0.85,
        size=11, color=BODY_TEXT)

# Venn-style 3-circle callouts (simplified as cards)
circles = [
    ("Right People", "Those that need\nsupport and help\nto change."),
    ("Right Program", "Enrolment and\noutcome criteria."),
    ("Right Time",    "Committed and\nready to change."),
]
for i, (title, desc) in enumerate(circles):
    lx = 8.2 + (i % 2) * 2.4
    ty = 1.4 + (i // 2) * 1.85 + (0 if i < 2 else 0)
    if i == 2:
        lx = 9.4
        ty = 3.5
    rect(slide, lx, ty, 2.2, 1.65, fill=TEAL if i != 2 else TEAL_DARK)
    txt(slide, title, lx + 0.12, ty + 0.15, 1.96, 0.38,
        size=10, bold=True, color=WHITE)
    txt(slide, desc, lx + 0.12, ty + 0.55, 1.96, 0.98,
        size=9, color=WHITE)

txt(slide, "Maximum\nROI",
    9.5, 2.45, 1.6, 0.8,
    size=16, bold=True, color=NAVY, align=PP_ALIGN.CENTER)

footer(slide, 11)


# ── SLIDE 12 — HOW IT WORKS (5 STEPS) ───────────────────────────────────────
slide = prs.slides.add_slide(BLANK)
white_bg(slide)

# left photo column
photo_placeholder(slide, 0.3, 0.3, 3.5, 6.7, "[ Partner meeting photo ]")

txt(slide, "How it works",
    4.2, 0.38, 8.7, 0.6,
    size=24, bold=True, color=DARK_TEXT)
txt(slide, "We work with you to build the right solution for your members and patients.",
    4.2, 0.98, 8.5, 0.38,
    size=11, color=TEAL)

# tick summary
for i, item in enumerate(["\u2713 Your success is our success",
                           "\u2713 Lower costs",
                           "\u2713 Increased outcomes"]):
    txt(slide, item, 4.2 + i * 2.85, 1.5, 2.75, 0.35,
        size=10, bold=True, color=TEAL)

steps = [
    ("Step 1: Build",
     "We help you choose a core solution(s) that will meet your business objectives, in line with your budget."),
    ("Step 2: Assess",
     "Based on your budget, desired outcomes and optional extras, we create options to walk you through."),
    ("Step 3: Sign",
     "Once a solution has been agreed, we move to contract stage and begin building the launch plan."),
    ("Step 4: Launch Plan",
     "We introduce teams, create supporting assets including marketing assistance, best practices and staff training."),
    ("Step 5: Review",
     "Monthly Catch-Ups with your dedicated Account Manager to review engagement metrics. Cohort reporting offered quarterly and annually.*"),
]
step_w = 1.8
for i, (title, body) in enumerate(steps):
    step_card(slide, 4.1 + i * 1.84, 2.05, step_w, title, "", body)
    # override the card header with smaller title
    rect(slide, 4.1 + i * 1.84, 2.05, step_w, 0.38, fill=TEAL)
    txt(slide, title, 4.15 + i * 1.84, 2.1, step_w - 0.1, 0.34,
        size=8, bold=True, color=WHITE)

txt(slide, "*Cohort reporting offered quarterly ($1,800) and annually ($2,100). Free for cohorts of 201+.",
    4.2, 6.6, 8.9, 0.32,
    size=7.5, color=MID_GREY, italic=True)

footer(slide, 12)


# ── SLIDE 13 — PROGRAM OPTIONS (CDMP) ───────────────────────────────────────
slide = prs.slides.add_slide(BLANK)
white_bg(slide)

txt(slide, "Total Wellbeing Lifestyle Plan — CDMP",
    0.55, 0.38, 9.5, 0.6,
    size=24, bold=True, color=DARK_TEXT)
txt(slide, "PHI-compliant Chronic Disease Management Program — centred around your members and patients.",
    0.55, 0.98, 9.5, 0.42,
    size=12, color=TEAL)

# tier table headers
headers = ["Program Tier", "Duration", "Health Coaching", "Fast Start Option", "Cost / Member*"]
col_w   = [2.8, 1.6, 2.0, 2.0, 2.6]
col_x   = [0.45]
for w in col_w[:-1]:
    col_x.append(col_x[-1] + w + 0.06)

header_y = 1.62
row_h    = 0.56

for ci, (hdr, cx, cw) in enumerate(zip(headers, col_x, col_w)):
    rect(slide, cx, header_y, cw, row_h - 0.06, fill=NAVY)
    txt(slide, hdr, cx + 0.1, header_y + 0.1, cw - 0.2, row_h - 0.25,
        size=10, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

tiers = [
    ["Entry — Prevention",     "12 weeks", "3 HCS",  "Optional", "From $595"],
    ["Standard — CDMP",        "16 weeks", "7–10 HCS","Optional", "From $866"],
    ["Comprehensive — CDMP",   "24 weeks", "10–15 HCS","Optional","From $1,125"],
]
row_colors = [LIGHT_BG, WHITE, LIGHT_BG]
for ri, (row_data, bg) in enumerate(zip(tiers, row_colors)):
    ry = header_y + row_h + ri * row_h
    for ci, (cell, cx, cw) in enumerate(zip(row_data, col_x, col_w)):
        rect(slide, cx, ry, cw, row_h - 0.06, fill=bg)
        txt(slide, cell, cx + 0.1, ry + 0.1, cw - 0.2, row_h - 0.22,
            size=10,
            color=TEAL if ci == 0 else BODY_TEXT,
            bold=(ci == 0),
            align=PP_ALIGN.CENTER if ci != 0 else PP_ALIGN.LEFT)

# eligibility note
txt(slide, "Eligibility criteria:",
    0.55, 3.56, 9.0, 0.38,
    size=11, bold=True, color=DARK_TEXT)

elig = [
    "BMI \u2265 28 kg/m\u00b2, AND",
    "One or more chronic conditions: Cardiovascular disease, Osteoarthritis, Type 2 Diabetes, "
    "Obstructive Sleep Apnoea, Respiratory Conditions (COPD/Asthma). OR",
    "Is at high risk: lifestyle risk factors + biomedical risk factors + family history of chronic disease.",
]
for i, item in enumerate(elig):
    txt(slide, f"\u2022  {item}", 0.55, 4.0 + i * 0.56, 8.8, 0.52,
        size=10, color=BODY_TEXT)

txt(slide, "*Pricing shown excludes GST. Volume discounts apply (30–45% off fixed costs for 100–300+ registrations). See pricing schedule for full detail.",
    0.55, 6.6, 12.3, 0.35,
    size=7.5, color=MID_GREY, italic=True)

photo_placeholder(slide, 10.2, 1.55, 2.8, 5.1, "[ Member photo ]")
footer(slide, 13)


# ── SLIDE 14 — SECTION DIVIDER: MORE INFORMATION ────────────────────────────
slide = prs.slides.add_slide(BLANK)
section_divider_white(
    slide, 14,
    "Digital Wellness\u00ae",
    "More\nInformation",
    "\u2713 Health Coaching Sessions   \u2713 Fast Start   \u2713 Reporting")


# ── SLIDE 15 — HEALTH COACHING SESSIONS ─────────────────────────────────────
slide = prs.slides.add_slide(BLANK)
white_bg(slide)

txt(slide, "Health Coaching Sessions",
    0.55, 0.38, 9.5, 0.6,
    size=24, bold=True, color=DARK_TEXT)
txt(slide, "Coaching support that drives member wellbeing and sustainable results.",
    0.55, 0.98, 9.5, 0.4,
    size=12, color=TEAL)

intro = ("Our Total Wellbeing Health Coaches are Accredited Practising Dietitians (APDs) and members "
         "of Dietitians Australia. They provide evidence-based support across nutrition, behaviour change "
         "and lifestyle planning — helping members set realistic goals, stay accountable and build lasting "
         "healthy habits.")
txt(slide, intro,
    0.55, 1.55, 7.8, 0.95,
    size=11, color=BODY_TEXT)

txt(slide, "What members can expect from a coaching session:",
    0.55, 2.65, 8.0, 0.38,
    size=11, bold=True, color=DARK_TEXT)

coaching_items = [
    "Personalised weight-loss and nutrition advice",
    "Progress review and recommendations to improve outcomes",
    "Guidance on getting the most from the CSIRO Total Wellbeing Diet platform",
    "Answers to program and meal-plan questions",
    "Motivation, accountability and practical strategies to stay on track",
    "Support for overcoming challenges and developing sustainable habits",
]
for i, item in enumerate(coaching_items):
    col = i % 2
    row = i // 2
    lx = 0.55 + col * 6.3
    ty = 3.1 + row * 0.62
    txt(slide, f"\u2713  {item}", lx, ty, 5.9, 0.55,
        size=10, color=BODY_TEXT)

photo_placeholder(slide, 9.8, 1.4, 3.2, 5.1, "[ Coaching session photo ]")
footer(slide, 15)


# ── SLIDE 16 — FAST START ────────────────────────────────────────────────────
slide = prs.slides.add_slide(BLANK)
white_bg(slide)

txt(slide, "Fast Start Meal Replacements",
    0.55, 0.38, 9.5, 0.6,
    size=24, bold=True, color=DARK_TEXT)
txt(slide, "Convenient nutrition on the go — boosting early weight loss and building confidence.",
    0.55, 0.98, 9.5, 0.4,
    size=12, color=TEAL)

# 4 feature cards
features = [
    ("Kickstarts Weight Loss",
     "Using meal replacements for breakfast and lunch in the first 3 weeks provides a "
     "boost in early weight loss, increasing confidence and motivation."),
    ("Phased Transition",
     "Weeks 1–3: two Fast Start Shakes per day. Weeks 4+: gradual transition to the "
     "Total Wellbeing Diet whole-food plan, with the flexibility to continue using "
     "shakes as needed."),
    ("Proven Results",
     "98% of participants achieve weight loss. 75% achieve clinically significant "
     "reductions (>5%). On average, participants experience a 7.4% reduction in body "
     "weight within 12 weeks."),
    ("Adaptable to Any Program",
     "Fast Start Shakes can be seamlessly integrated into any Total Wellbeing Diet "
     "program — available in vanilla, strawberry, coffee and chocolate. The 3-Week "
     "Starter Pack includes 42 shakes and a shaker."),
]
for i, (title, desc) in enumerate(features):
    col = i % 2
    row = i // 2
    lx = 0.45 + col * 4.85
    ty = 1.58 + row * 2.2
    rect(slide, lx, ty, 4.55, 2.0, fill=LIGHT_BG)
    rect(slide, lx, ty, 0.1, 2.0, fill=TEAL)
    txt(slide, title, lx + 0.22, ty + 0.14, 4.18, 0.4,
        size=11, bold=True, color=TEAL)
    txt(slide, desc, lx + 0.22, ty + 0.6, 4.18, 1.3,
        size=10, color=BODY_TEXT)

# fast start stats bar
rect(slide, 0.45, 6.0, 9.45, 0.68, fill=NAVY)
fs_stats = [("98%", "experience weight loss"), ("75%", "clinically significant loss"),
            ("7.4%", "avg body weight loss"), ("80%", "better craving control at 3 weeks")]
for i, (val, lbl) in enumerate(fs_stats):
    lx = 0.65 + i * 2.38
    txt(slide, val, lx, 6.06, 1.2, 0.34,
        size=14, bold=True, color=TEAL)
    txt(slide, lbl, lx + 1.22, 6.1, 1.1, 0.3,
        size=8, color=WHITE)

photo_placeholder(slide, 10.1, 1.4, 3.0, 5.3, "[ Fast Start product photo ]")
footer(slide, 16)


# ── SLIDE 17 — CLOSING / NEXT STEPS ─────────────────────────────────────────
slide = prs.slides.add_slide(BLANK)
navy_bg(slide)
dot_pattern(slide, 0.3, 0.3, cols=6, rows=4, spacing=0.25,
            r=0.045, color=RGBColor(0x25, 0x3C, 0x5E))

rect(slide, 0, 0, 0.22, 7.5, fill=TEAL)

txt(slide, "Making better health possible —",
    0.6, 1.1, 11.5, 0.7,
    size=30, bold=False, color=WHITE)
txt(slide, "for everyone, for life.",
    0.6, 1.78, 11.5, 0.7,
    size=30, bold=True, color=TEAL)

rect(slide, 0.6, 2.65, 11.5, 0.04, fill=TEAL)

takeaways = [
    ("Science is our moat.",
     "CSIRO and Mayo Clinic are the most trusted health research brands in their markets. No competitor can replicate this depth of peer-reviewed authority."),
    ("Prevention pays.",
     "Every dollar invested in upstream risk identification and chronic disease management reduces future claims, hospitalisation and workforce absenteeism."),
    ("We walk alongside you.",
     "From program design to launch, marketing support, staff training and ongoing reporting — your success is our success."),
]
for i, (bold_txt, rest) in enumerate(takeaways):
    lx = 0.6 + i * 4.15
    rect(slide, lx, 2.85, 3.9, 3.5, fill=RGBColor(0x25, 0x3C, 0x5E))
    rect(slide, lx, 2.85, 3.9, 0.07,
         fill=TEAL if i == 0 else (RGBColor(0x00, 0x9A, 0xA0) if i == 1 else TEAL_DARK))
    txt(slide, bold_txt, lx + 0.2, 3.02, 3.5, 0.45,
        size=13, bold=True, color=WHITE)
    txt(slide, rest, lx + 0.2, 3.52, 3.5, 1.7,
        size=10, color=MID_GREY)

txt(slide, "Digital Wellness\u00ae   |   www.digitalwellness.com   |   www.totalwellbeingdiet.com",
    0.6, 6.9, 12.0, 0.32,
    size=9, color=MID_GREY, align=PP_ALIGN.CENTER)

footer(slide, 17)


# ── Save ─────────────────────────────────────────────────────────────────────
out = "/home/user/ClaudeCodeTest/reports/enterprise-proposal-FY26.pptx"
prs.save(out)
print(f"Saved: {out}")
